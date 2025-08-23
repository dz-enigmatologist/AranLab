import os
from openpyxl import Workbook, load_workbook

def write_headers_to_worksheet(worksheet):
    headers = [
        'time(mins)', 'delta Rct-a', 'normalized Rct_a','delta Rct-d','normalized Rct_d', 'Cp1', 'Ph1', 
        'Slope 1', 'Slope 2', 'Slope 3', 'Slope 4', 'Slope 5', 
        'Angle', 'Cp_exp-a', 'Cp_exp-b', 'Ph_slope', 'Ph_peak', 
        'Area Cp', 'Area Ph', 'Area Slope', 'Area Rs-direct', 'Area Rs-Para',
        '','','linear_eq_m','linear_eq_b','Rs','delta Rct-i','Q','n'
    ]
    for col_num, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=col_num, value=header)


def create_chips_folder_and_workbooks(folder_path=None):
    """Create all chip workbooks only if they do not already exist"""
    if folder_path is None:
        folder_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Processed Chips for PCAI1ia")
    os.makedirs(folder_path, exist_ok=True)

    workbook_names = [
        'Chip 21', 'Chip 22', 'Chip 23', 'Chip 26', 'Chip 27',
        'Chip 32', 'Chip 33', 'Chip 35', 'Chip 36', 'Chip 37',
        'Chip 39', 'Chip 40', 'Chip 41', 'Chip 43', 'Chip 44',
        'Chip 45', 'Chip 46', 'Chip 47', 'Chip 48', 'Chip 52',
        'Chip 53', 'Chip 54', 'Chip 55', 'Chip 56', 'Chip 57',
        'Chip 59'
    ]
    worksheet_names = ["0pM_asso", "0pM_disso", "100pM_asso", "100pM_disso",
                       "1nM_asso", "1nM_disso", "10nM_asso", "10nM_disso",
                       "100nM_asso", "100nM_disso"]

    for name in workbook_names:
        file_path = os.path.join(folder_path, f"{name}.xlsx")
        if os.path.exists(file_path):
            print(f"Workbook '{name}' already exists. Skipping creation.")
            continue

        try:
            wb = Workbook()
            if wb.active:
                wb.remove(wb.active)
            for sheet_name in worksheet_names:
                ws = wb.create_sheet(title=sheet_name)
                write_headers_to_worksheet(ws)
            wb.save(file_path)
            print(f"Created workbook '{name}' at {file_path}")
        except Exception as e:
            print(f"Error creating workbook {name}: {e}")


def create_single_chip_workbook(chip_name, folder_path=None):
    """Create a single chip workbook only if it does not already exist"""
    if folder_path is None:
        folder_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Processed Chips for PCAI1ia")
    os.makedirs(folder_path, exist_ok=True)

    file_path = os.path.join(folder_path, f"{chip_name}.xlsx")
    if os.path.exists(file_path):
        print(f"Workbook '{chip_name}' already exists. Skipping creation.")
        return

    worksheet_names = ["0pM_asso", "0pM_disso", "100pM_asso", "100pM_disso",
                       "1nM_asso", "1nM_disso", "10nM_asso", "10nM_disso",
                       "100nM_asso", "100nM_disso"]
    try:
        wb = Workbook()
        if wb.active:
            wb.remove(wb.active)
        for sheet_name in worksheet_names:
            ws = wb.create_sheet(title=sheet_name)
            write_headers_to_worksheet(ws)
        wb.save(file_path)
        print(f"Created workbook '{chip_name}' at {file_path}")
    except Exception as e:
        print(f"Error creating workbook {chip_name}: {e}")


def append_rows_to_sheet(file_path, sheet_name, rows):
    """Append rows of data to an existing sheet in an Excel workbook."""
    if not os.path.exists(file_path):
        print(f"File '{file_path}' does not exist. Cannot append.")
        return

    wb = load_workbook(file_path)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        write_headers_to_worksheet(ws)
    else:
        ws = wb[sheet_name]

    for row in rows:
        ws.append(row)

    wb.save(file_path)
